from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse

from pydantic import BaseModel

from io import BytesIO

from docx import Document

from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer
)
from reportlab.lib.styles import getSampleStyleSheet


from models import (
    TranscriptRequest,
    RequirementsRequest
)

from ai_service import (
    analyze_transcript,
    generate_requirements,
    evaluate_requirements
)


# ===========================================
# FASTAPI APP
# ===========================================

app = FastAPI(
    title="AI Requirement Analyzer API"
)


# ===========================================
# CORS
# ===========================================

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ===========================================
# EXPORT REQUEST MODEL
# ===========================================

class ExportRequest(BaseModel):

    transcript: str

    analysis: str = ""

    clarifications: str = ""

    requirements: str = ""

    evaluation: str = ""


# ===========================================
# HELPER: CREATE REPORT TEXT
# ===========================================

def create_report_text(request: ExportRequest):

    return f"""
AI REQUIREMENT ANALYZER REPORT
========================================


MEETING TRANSCRIPT
----------------------------------------

{request.transcript}


AI ANALYSIS
----------------------------------------

{request.analysis}


CLARIFICATION ANSWERS
----------------------------------------

{request.clarifications}


GENERATED REQUIREMENTS
----------------------------------------

{request.requirements}


QUALITY EVALUATION
----------------------------------------

{request.evaluation}

"""


# ===========================================
# HOME
# ===========================================

@app.get("/")
def home():

    return {
        "message":
            "AI Requirement Analyzer Backend is Running!"
    }


# ===========================================
# HEALTH CHECK
# ===========================================

@app.get("/health")
def health_check():

    return {
        "status": "healthy"
    }


# ===========================================
# 1. ANALYZE TRANSCRIPT
# ===========================================

@app.post("/analyze-transcript")
def analyze(request: TranscriptRequest):

    try:

        result = analyze_transcript(
            request.transcript
        )

        return {
            "analysis": result
        }

    except Exception as e:

        raise HTTPException(
            status_code=500,
            detail=str(e)
        )


# ===========================================
# 2. GENERATE REQUIREMENTS
# ===========================================

@app.post("/generate-requirements")
def generate(request: RequirementsRequest):

    try:

        clarifications_text = ""

        for item in request.clarifications:

            clarifications_text += (

                f"Question: {item.question}\n"

                f"Answer: {item.answer}\n\n"

            )


        result = generate_requirements(

            transcript=request.transcript,

            clarifications=clarifications_text

        )


        return {

            "requirements": result

        }


    except Exception as e:

        raise HTTPException(

            status_code=500,

            detail=str(e)

        )


# ===========================================
# 3. EVALUATE REQUIREMENTS
# ===========================================

@app.post("/evaluate")
def evaluate(request: RequirementsRequest):

    try:

        # -----------------------------------
        # REQUIREMENTS WITHOUT CLARIFICATION
        # -----------------------------------

        before = generate_requirements(

            transcript=request.transcript,

            clarifications=""

        )


        # -----------------------------------
        # CONVERT CLARIFICATIONS TO TEXT
        # -----------------------------------

        clarifications_text = ""

        for item in request.clarifications:

            clarifications_text += (

                f"Question: {item.question}\n"

                f"Answer: {item.answer}\n\n"

            )


        # -----------------------------------
        # REQUIREMENTS WITH CLARIFICATION
        # -----------------------------------

        after = generate_requirements(

            transcript=request.transcript,

            clarifications=clarifications_text

        )


        # -----------------------------------
        # AI EVALUATION
        # -----------------------------------

        evaluation = evaluate_requirements(

            before=before,

            after=after

        )


        return {

            "without_clarification": before,

            "with_clarification": after,

            "evaluation": evaluation

        }


    except Exception as e:

        raise HTTPException(

            status_code=500,

            detail=str(e)

        )


# ===========================================
# 4. EXPORT TXT
# ===========================================

@app.post("/export/txt")
def export_txt(request: ExportRequest):

    try:

        report = create_report_text(request)

        file = BytesIO()

        file.write(
            report.encode("utf-8")
        )

        file.seek(0)


        return StreamingResponse(

            file,

            media_type="text/plain",

            headers={

                "Content-Disposition":
                    "attachment; filename=AI_Requirement_Analyzer_Report.txt"

            }

        )


    except Exception as e:

        raise HTTPException(

            status_code=500,

            detail=str(e)

        )


# ===========================================
# 5. EXPORT DOCX
# ===========================================

@app.post("/export/docx")
def export_docx(request: ExportRequest):

    try:

        document = Document()


        document.add_heading(

            "AI Requirement Analyzer Report",

            level=0

        )


        sections = [

            (
                "Meeting Transcript",

                request.transcript

            ),

            (
                "AI Analysis",

                request.analysis

            ),

            (
                "Clarification Answers",

                request.clarifications

            ),

            (
                "Generated Requirements",

                request.requirements

            ),

            (
                "Quality Evaluation",

                request.evaluation

            )

        ]


        for title, content in sections:

            document.add_heading(

                title,

                level=1

            )


            document.add_paragraph(

                content

            )


        file = BytesIO()

        document.save(file)

        file.seek(0)


        return StreamingResponse(

            file,

            media_type=
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",

            headers={

                "Content-Disposition":
                    "attachment; filename=AI_Requirement_Analyzer_Report.docx"

            }

        )


    except Exception as e:

        raise HTTPException(

            status_code=500,

            detail=str(e)

        )


# ===========================================
# 6. EXPORT PDF
# ===========================================

@app.post("/export/pdf")
def export_pdf(request: ExportRequest):

    try:

        file = BytesIO()


        document = SimpleDocTemplate(

            file,

            pagesize=A4,

            rightMargin=40,

            leftMargin=40,

            topMargin=40,

            bottomMargin=40

        )


        styles = getSampleStyleSheet()


        story = []


        story.append(

            Paragraph(

                "AI Requirement Analyzer Report",

                styles["Title"]

            )

        )


        story.append(

            Spacer(

                1,

                20

            )

        )


        sections = [

            (
                "Meeting Transcript",

                request.transcript

            ),

            (
                "AI Analysis",

                request.analysis

            ),

            (
                "Clarification Answers",

                request.clarifications

            ),

            (
                "Generated Requirements",

                request.requirements

            ),

            (
                "Quality Evaluation",

                request.evaluation

            )

        ]


        for title, content in sections:

            story.append(

                Paragraph(

                    title,

                    styles["Heading2"]

                )

            )


            story.append(

                Spacer(

                    1,

                    8

                )

            )


            formatted_content = (

                content

                .replace("&", "&amp;")

                .replace("<", "&lt;")

                .replace(">", "&gt;")

                .replace("\n", "<br/>")

            )


            story.append(

                Paragraph(

                    formatted_content,

                    styles["BodyText"]

                )

            )


            story.append(

                Spacer(

                    1,

                    18

                )

            )


        document.build(story)

        file.seek(0)


        return StreamingResponse(

            file,

            media_type="application/pdf",

            headers={

                "Content-Disposition":
                    "attachment; filename=AI_Requirement_Analyzer_Report.pdf"

            }

        )


    except Exception as e:

        raise HTTPException(

            status_code=500,

            detail=str(e)

        )